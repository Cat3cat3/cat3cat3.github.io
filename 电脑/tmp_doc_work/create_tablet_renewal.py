from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path("/Users/kevin/Documents/CATmaker/03-电脑租赁业务/成都市双流区酷虎科技培训学校有限公司")
TEMPLATE = ROOT / "电脑租赁续签协议.docx"
OUT = ROOT / "平板租赁续签协议.docx"


paragraphs = [
    "平板租赁续签协议",
    "甲方（出租方）：成都猫匠极客科技有限公司",
    "统一社会信用代码：91510100MAEPK5XC6E",
    "联系地址：四川省成都市双流区鲢鱼社区 11 幢 2 单元 CATmaker",
    "联系人 / 电话：吴文凯 18200456012",
    "乙方（承租方）：成都市双流区酷虎科技培训学校有限公司",
    "统一社会信用代码：91510116MACQB3XL2A",
    "联系地址：四川省成都市双流区东升街道万达广场 2 楼",
    "联系人 / 电话：高靖 13438285818",
    "",
    "续签基础信息",
    "原合同：双方于 2025 年 06 月 27 日签订《租赁技术服务合同》（订单号：N20250627000001），原租期至 2026 年 06 月 27 日届满。",
    "续租设备：准新 9.7 寸 iPad 6，共 2 台，设备配置、序列号、数量及设备价值均按原合同不变。",
    "续租期限：自 2026 年 06 月 28 日起至 2027 年 06 月 27 日止，共计 1 年。",
    "续租租金：每台每月人民币 48 元（大写：肆拾捌元整），共 2 台，每月合计人民币 96 元（大写：玖拾陆元整），按自然月结算。",
    "",
    "核心约定",
    "本协议为原《租赁技术服务合同》的续签补充协议，与原合同具有同等法律效力。",
    "本协议仅对续租期间的租金标准进行调整，原合同其他全部条款（服务、维修、违约责任、归还、争议解决等）均继续有效、不作变更。",
    "续租期间仍执行：租期 1 年，相关规则按原合同第六条第 3 款执行。",
    "续租期间押金：0 元（延续原合同免押金）。",
    "付款方式：先付后用，每月 27 日前支付下月租金，支付账户同原合同：",
    "户名：吴文凯",
    "账号：6214 8328 1098 3531",
    "开户行：招商银行成都龙湖三千支行",
    "发票：甲方按乙方实际支付金额开具电子发票。",
    "",
    "其他",
    "本协议自双方签字盖章之日起生效，一式两份，甲乙双方各执一份，具有同等法律效力。",
    "续租期满，双方可另行协商续签；未续签且乙方继续使用设备的，本协议约定的租金标准自动顺延。",
    "",
    "甲方（盖章 / 签字）：__________日期：______年____月____日",
    "乙方（盖章 / 签字）：__________日期：______年____月____日",
]


def set_paragraph_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def main():
    doc = Document(TEMPLATE)

    for idx, text in enumerate(paragraphs):
        if idx < len(doc.paragraphs):
            set_paragraph_text(doc.paragraphs[idx], text)
        else:
            source = doc.paragraphs[-1]
            new_p = deepcopy(source._p)
            source._p.addnext(new_p)
            new_para = doc.paragraphs[-1]
            set_paragraph_text(new_para, text)

    for extra in doc.paragraphs[len(paragraphs):]:
        set_paragraph_text(extra, "")

    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for idx in (10, 16, 27):
        if idx < len(doc.paragraphs):
            doc.paragraphs[idx].alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.core_properties.title = "平板租赁续签协议"
    doc.core_properties.subject = "租赁续签协议"
    doc.core_properties.keywords = "平板租赁;续签协议;iPad;酷虎科技培训学校"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
