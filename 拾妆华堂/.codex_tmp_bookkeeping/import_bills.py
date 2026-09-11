import csv
import hashlib
import json
import os
import re
import shutil
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None


BASE = Path("/Users/kevin/Documents/CATmaker/01-拾妆华堂")
BOOK_DIR = BASE / "化妆工作室账簿"
WORKBOOK = BOOK_DIR / "化妆工作室账簿.xlsx"
INTAKE = BOOK_DIR / "账单导入" / "2026-07-05"
EXTRACTED = INTAKE / "extracted"
SCAN_START = date(2026, 5, 19)

HEADERS = [
    "日期", "平台", "类型", "金额", "分类", "摘要", "来源文件", "置信度", "备注",
    "确认状态", "来源冲突", "记录ID", "处理批次", "原始文本",
]

ZIP_SPECS = [
    ("微信", "微信支付账单流水文件*.zip", "WECHAT_ZIP_PWD"),
    ("支付宝", "支付宝交易明细*.zip", "ALIPAY_ZIP_PWD"),
    ("招商银行", "招商银行交易流水*.zip", "CMB_ZIP_PWD"),
]

BUSINESS_HINTS = [
    "工作室", "化妆", "妆", "美妆", "影楼", "服装", "道具", "灯", "灯棒", "神牛",
    "lr500", "口红", "粉底", "眼影", "睫毛", "假发", "发饰", "头饰", "汉服",
    "布料", "镜", "收纳", "打印", "相册", "摄影", "拍摄", "快递", "运费", "淘宝",
    "天猫", "拼多多", "闲鱼", "抖音", "美团", "1688",
]
PERSONAL_HINTS = [
    "中国电信", "话费", "地铁", "公交", "滴滴", "瑞幸", "星巴克", "麦当劳", "肯德基",
    "大米先生", "可口可乐", "便利店", "超市", "外卖", "饿了么", "电影", "游戏",
]


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(BOOK_DIR))
    except ValueError:
        return str(path)


def clean_text(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_amount(value):
    if value is None:
        return None
    text = str(value).replace(",", "").replace("￥", "").replace("¥", "").strip()
    m = re.search(r"-?\d+(?:\.\d+)?", text)
    return round(float(m.group(0)), 2) if m else None


def parse_dt(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    for pattern in (
        r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})",
        r"(\d{4})(\d{2})(\d{2})",
    ):
        m = re.search(pattern, text)
        if m:
            try:
                return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                return None
    return None


def infer_type(amount, text):
    lower = text.lower()
    if any(k in text for k in ["退款", "退回", "已全额退款"]):
        return "退款"
    if any(k in text for k in ["收入", "收款", "转入", "贷方", "入账", "代发", "红包收入"]):
        return "收入"
    if any(k in text for k in ["支出", "付款", "消费", "借方", "转出", "提现", "还款"]):
        return "成本" if any(h.lower() in lower for h in BUSINESS_HINTS) else "支出"
    if amount is not None and amount < 0:
        return "支出"
    return "待判断"


def infer_note_and_conf(text, tx_type):
    lower = text.lower()
    business = any(h.lower() in lower for h in BUSINESS_HINTS)
    personal = any(h.lower() in lower for h in PERSONAL_HINTS)
    if business and not personal:
        note = "疑似公司交易，请确认用途"
        conf = "中"
    elif personal and not business:
        note = "疑似个人交易，请确认是否排除"
        conf = "中"
    else:
        note = "个人/公司归属不明确，请确认"
        conf = "低" if tx_type == "待判断" else "中"
    return note, conf


def record_id(tx_date, platform, amount, summary):
    raw = f"{tx_date}|{platform}|{amount:.2f}|{summary}".lower()
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def extract_archives():
    EXTRACTED.mkdir(exist_ok=True)
    extracted = []
    for platform, glob_pat, env_name in ZIP_SPECS:
        matches = list(INTAKE.glob(glob_pat))
        if not matches:
            raise FileNotFoundError(f"未找到 {platform} 压缩包")
        pwd = os.environ.get(env_name)
        if not pwd:
            raise RuntimeError(f"缺少 {platform} 密码环境变量")
        out_dir = EXTRACTED / platform
        out_dir.mkdir(exist_ok=True)
        with zipfile.ZipFile(matches[0]) as zf:
            zf.extractall(out_dir, pwd=pwd.encode("utf-8"))
        for p in out_dir.rglob("*"):
            if p.is_file():
                extracted.append(p)
    return extracted


def parse_wechat_xlsx(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    headers = []
    for i, row in enumerate(rows):
        vals = [clean_text(v) for v in row]
        if "交易时间" in vals and "金额(元)" in vals:
            header_idx = i
            headers = vals
            break
    if header_idx is None:
        return []
    idx = {h: n for n, h in enumerate(headers) if h}
    out = []
    for row in rows[header_idx + 1:]:
        vals = [clean_text(v) for v in row]
        if not any(vals):
            continue
        d = parse_dt(vals[idx.get("交易时间", 0)])
        amount = parse_amount(vals[idx.get("金额(元)", -1)])
        if not d or amount is None:
            continue
        direction = vals[idx.get("收/支", -1)] if "收/支" in idx else ""
        merchant = vals[idx.get("交易对方", -1)] if "交易对方" in idx else ""
        product = vals[idx.get("商品", -1)] if "商品" in idx else ""
        status = vals[idx.get("当前状态", -1)] if "当前状态" in idx else ""
        summary = " / ".join([x for x in [merchant, product, status] if x])[:120]
        raw = " | ".join(vals)
        signed_amount = amount if direction == "收入" else -amount if direction == "支出" else amount
        out.append(make_candidate(d, "微信", signed_amount, summary, raw, path))
    return out


def read_text_file(path: Path):
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030", "gbk", "utf-16"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace"), "utf-8-replace"


def parse_alipay_csv(path: Path):
    text, _ = read_text_file(path)
    lines = text.splitlines()
    header_idx = None
    for i, line in enumerate(lines):
        if "交易创建时间" in line and ("金额" in line or "收入" in line or "支出" in line):
            header_idx = i
            break
    if header_idx is None:
        return []
    sample = "\n".join(lines[header_idx:header_idx + 5])
    dialect = csv.Sniffer().sniff(sample, delimiters=",\t")
    reader = csv.DictReader(lines[header_idx:], dialect=dialect)
    out = []
    for row in reader:
        d = parse_dt(row.get("交易创建时间") or row.get("付款时间") or row.get("最近修改时间"))
        if not d:
            continue
        income = parse_amount(row.get("收入（+元）") or row.get("收入(+元)") or row.get("收入"))
        expense = parse_amount(row.get("支出（-元）") or row.get("支出(-元)") or row.get("支出"))
        amount = income if income is not None else -expense if expense is not None else parse_amount(row.get("金额"))
        if amount is None:
            continue
        summary = " / ".join(
            clean_text(row.get(k)) for k in ["交易对方", "商品说明", "交易状态", "交易分类"]
            if clean_text(row.get(k))
        )[:120]
        raw = " | ".join(f"{clean_text(k)}={clean_text(v)}" for k, v in row.items())
        out.append(make_candidate(d, "支付宝", amount, summary, raw, path))
    return out


def parse_cmb_pdf(path: Path):
    if PdfReader is None:
        return []
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    out = []
    for line in text.splitlines():
        line = clean_text(line)
        if not re.search(r"2026[-/.年]\d{1,2}[-/.月]\d{1,2}", line):
            continue
        d = parse_dt(line)
        amounts = [parse_amount(x) for x in re.findall(r"-?\d[\d,]*\.\d{2}", line)]
        amounts = [x for x in amounts if x is not None]
        if not d or not amounts:
            continue
        amount = amounts[0]
        if any(k in line for k in ["支出", "借方", "消费", "转出"]):
            amount = -abs(amount)
        elif any(k in line for k in ["收入", "贷方", "转入", "入账"]):
            amount = abs(amount)
        summary = re.sub(r"\d[\d,]*\.\d{2}", "", line)
        summary = re.sub(r"2026[-/.年]\d{1,2}[-/.月]\d{1,2}[日]?", "", summary).strip()[:120]
        out.append(make_candidate(d, "招商银行", amount, summary, line, path))
    return out


def make_candidate(d, platform, amount, summary, raw, source):
    tx_type = infer_type(amount, raw + " " + summary)
    note, conf = infer_note_and_conf(raw + " " + summary, tx_type)
    category = "待分类"
    if tx_type in ("成本", "支出"):
        category = "购买费用" if "疑似公司" in note else "待分类"
    rid = record_id(d.isoformat(), platform, abs(amount), summary)
    return {
        "日期": d,
        "平台": platform,
        "类型": tx_type,
        "金额": abs(amount),
        "分类": category,
        "摘要": summary or "未识别摘要",
        "来源文件": rel(source),
        "置信度": conf,
        "备注": note,
        "确认状态": "待确认",
        "来源冲突": "",
        "记录ID": rid,
        "处理批次": datetime.now().strftime("%Y%m%d-%H%M%S-auto-bills"),
        "原始文本": raw[:500],
    }


def parse_all(files):
    candidates = []
    for p in files:
        suffix = p.suffix.lower()
        if suffix == ".xlsx":
            candidates.extend(parse_wechat_xlsx(p))
        elif suffix == ".csv":
            candidates.extend(parse_alipay_csv(p))
        elif suffix == ".pdf":
            candidates.extend(parse_cmb_pdf(p))
    return candidates


def existing_keys(ws):
    ids = set()
    signatures = set()
    header = [cell.value for cell in ws[1]]
    h = {name: i + 1 for i, name in enumerate(header)}
    for row in range(2, ws.max_row + 1):
        rid = ws.cell(row, h.get("记录ID", 12)).value
        if rid:
            ids.add(str(rid))
        d = ws.cell(row, h.get("日期", 1)).value
        platform = clean_text(ws.cell(row, h.get("平台", 2)).value)
        amount = ws.cell(row, h.get("金额", 4)).value
        summary = clean_text(ws.cell(row, h.get("摘要", 6)).value)
        pd = parse_dt(d)
        pa = parse_amount(amount)
        if pd and pa is not None:
            signatures.add((pd.isoformat(), platform, round(abs(pa), 2), summary[:32]))
    return ids, signatures


def formal_signatures(wb):
    sigs = set()
    if "收支财务报表" not in wb.sheetnames:
        return sigs
    ws = wb["收支财务报表"]
    for row in range(2, ws.max_row + 1):
        d = parse_dt(ws.cell(row, 2).value)
        summary = clean_text(ws.cell(row, 3).value)
        amount = parse_amount(ws.cell(row, 6).value)
        if d and amount is not None:
            sigs.add((d.isoformat(), round(abs(amount), 2), summary[:18]))
    return sigs


def update_workbook(candidates):
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = WORKBOOK.with_name(f"化妆工作室账簿.{stamp}.pre-auto-import.bak.xlsx")
    shutil.copy2(WORKBOOK, backup)
    wb = load_workbook(WORKBOOK)
    ws = wb["待确认"] if "待确认" in wb.sheetnames else wb.create_sheet("待确认")
    if ws.max_row == 1 and not ws.cell(1, 1).value:
        for col, h in enumerate(HEADERS, 1):
            ws.cell(1, col, h)
    else:
        existing_header = [ws.cell(1, col).value for col in range(1, len(HEADERS) + 1)]
        if existing_header != HEADERS:
            for col, h in enumerate(HEADERS, 1):
                ws.cell(1, col, h)

    ids, sigs = existing_keys(ws)
    formal = formal_signatures(wb)
    appended = []
    skipped = Counter()
    seen_batch = set()
    for c in sorted(candidates, key=lambda x: (x["日期"], x["平台"], x["金额"], x["摘要"])):
        if c["日期"] < SCAN_START:
            skipped["早于扫描起点"] += 1
            continue
        batch_sig = (c["日期"].isoformat(), c["平台"], round(c["金额"], 2), c["摘要"][:32])
        if c["记录ID"] in ids or batch_sig in sigs or batch_sig in seen_batch:
            skipped["重复"] += 1
            continue
        if any(c["日期"].isoformat() == fd and round(c["金额"], 2) == fa for fd, fa, _ in formal):
            c["来源冲突"] = "金额/日期与正式账簿可能重复"
            c["置信度"] = "低"
        row = ws.max_row + 1
        for col, h in enumerate(HEADERS, 1):
            val = c[h]
            ws.cell(row, col, val.isoformat() if isinstance(val, date) else val)
        appended.append(c)
        seen_batch.add(batch_sig)
    wb.save(WORKBOOK)
    return backup, appended, skipped


def main():
    files = extract_archives()
    candidates = parse_all(files)
    backup, appended, skipped = update_workbook(candidates)
    by_type = Counter(c["类型"] for c in appended)
    by_platform = Counter(c["平台"] for c in appended)
    sums = defaultdict(float)
    notes = Counter(c["备注"] for c in appended)
    for c in appended:
        sums[c["类型"]] += c["金额"]
    report = {
        "extracted_files": [rel(p) for p in files],
        "parsed_candidates": len(candidates),
        "appended": len(appended),
        "skipped": dict(skipped),
        "by_type": dict(by_type),
        "by_platform": dict(by_platform),
        "sums_by_type": {k: round(v, 2) for k, v in sums.items()},
        "notes": dict(notes),
        "backup": str(backup),
        "workbook": str(WORKBOOK),
        "scan_start": SCAN_START.isoformat(),
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
