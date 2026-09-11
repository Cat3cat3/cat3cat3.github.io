from pathlib import Path

from docx import Document


WORKSPACE = Path(__file__).resolve().parent


REPLACEMENTS = {
    "夜场直播化妆培训合作报价函（客户版）.docx": {
        "60,000元起 / 期": "80,000元起 / 期",
        "20,000元": "约26,667元",
        "60,000元": "80,000元",
        "68,000元": "88,000元",
        "76,000元": "96,000元",
    },
    "成都个人化妆培训费用市场对比参考（人力公司版）.docx": {
        "60,000元/期": "80,000元/期",
        "60,000元起/期": "80,000元起/期",
        "68,000元/期": "88,000元/期",
        "76,000元/期": "96,000元/期",
        "约6,000-7,500元/人": "约8,000-10,000元/人",
        "约5,231-6,182元/人": "约6,769-8,000元/人",
        "约5,067-5,429元/人": "约6,400-6,857元/人",
        "约5,067-7,500元/人": "约6,400-10,000元/人",
        "低于多数市场专项就业路径的中位价格": "接近市场专项就业路径的中低位价格",
        "低于或接近个人在成都报名基础妆+直播妆+夜场妆专项的常见区间": "处在个人报名基础妆+直播妆+夜场妆专项的常见价格区间内",
    },
}


def replace_in_paragraph(paragraph, mapping):
    for run in paragraph.runs:
        text = run.text
        for old, new in mapping.items():
            text = text.replace(old, new)
        run.text = text


def update_docx(path, mapping):
    doc = Document(path)
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, mapping)
    doc.save(path)


def main():
    for filename, mapping in REPLACEMENTS.items():
        update_docx(WORKSPACE / filename, mapping)
        print(filename)


if __name__ == "__main__":
    main()
